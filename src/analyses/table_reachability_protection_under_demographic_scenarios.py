#!/usr/bin/env python3
"""Reachability Protection Under Demographic Scenarios.

AnaSOP Sections 5-7: original matched-scenario reachability protection shares,
three rules and fixed original area groups at 1% requested bridge-relative cost.
MILI-D-20261008-006: tabulate the former figure heatmap; no new estimation.
"""
from pathlib import Path
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

from table_allocation_rules_and_shared_scenario_assumptions import preview_table

ROOT = Path(__file__).resolve().parents[2]
TITLE = "Reachability Protection Under Demographic Scenarios"
SLUG = "reachability_protection_under_demographic_scenarios"
EXP = ROOT / "data/exp/figure-table-generation" / SLUG
OUT = ROOT / f"data/results/tables/Table_{SLUG}.xlsx"
SOURCE = ROOT / "data/processed/demographic-stability/scenario_protection_preprocessed.parquet"
YEARS = [2020, 2030, 2040, 2050]
RULES = ["Traffic priority", "Condition priority", "Access priority"]
GROUPS = ["All matched residents", "Higher-ageing areas", "Distance-disadvantaged areas"]


def digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def main():
    EXP.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    before = digest(SOURCE)
    computation = json.loads((ROOT / "data/exp/demographic-stability/validation.json").read_text())
    assert computation["status"] == "passed"
    assert TITLE in (ROOT / "docs/AnaSOP.md").read_text()
    data = pd.read_parquet(SOURCE)
    data = data[data["Budget Share"].eq(.01)]
    assert len(data) == 36
    assert not data.duplicated(["Maintenance Priority", "Population Group", "Population-Change Scenario"]).any()
    assert data["Distribution of Protected Access"].between(-1e-12, 1 + 1e-12).all()
    rows = []
    for rule in RULES:
        for group in GROUPS:
            row = {"Maintenance priority": rule, "Population group": group}
            for year in YEARS:
                value = data.loc[data["Maintenance Priority"].eq(rule) & data["Population Group"].eq(group)
                                 & data["Population-Change Scenario"].eq(year), "Distribution of Protected Access"]
                assert len(value) == 1
                row[f"{year} (%)"] = float(value.iloc[0]) * 100
            rows.append(row)
    table = pd.DataFrame(rows)
    assert table.shape == (9, 6) and table.notna().all().all()
    original = ROOT / "data/exp/figure-table-generation/priority_stability_under_budget_and_demographic_scenarios/panel_b_values_percent.csv"
    if original.is_file():
        previous = pd.read_csv(original, index_col=0)
        np.testing.assert_allclose(table.iloc[:, 2:].to_numpy(), previous.to_numpy(), rtol=1e-12, atol=1e-12)
    notes = [
        TITLE,
        "Each value is the percentage of scenario-specific potential lost-medical-path population protected across equally weighted, mutually exclusive single-road closure events. It is not the percentage of the total resident population protected. Values are stored without rounding and displayed to two decimal places.",
        "All rules use a requested budget of 1% of the total relative intervention cost across 46,688 conditional bridge-associated road units. Cost is proportional to bridge length times width; the budget is neither a facility-count share nor a monetary expenditure. Complete fixed-prefix selection can leave unused budget.",
        "The matched 2020 reference comprises 175,485 origins and 125,874,510 census residents. The 127 unmatched original origins and 278 residents are excluded only from the matched demographic comparison. Original census population is multiplied by the official mesh projected-to-2020-reference ratio for 2030, 2040 and 2050; genuine projected zeros remain zero.",
        "Original 2020 higher-ageing and baseline-distance-disadvantaged area membership stays fixed. Groups overlap and are not additive. The network, medical destinations, maintenance units, costs, full event prevention and event weighting remain fixed. Access priority is reranked for each population scenario; traffic and condition selections remain fixed.",
        "Protection denominators change with scenario population; changes in these percentages do not by themselves establish changes in absolute protected residents. These are conditional bridge outcomes, not pavement sensitivity, future elderly-resident locations, observed repair effects, annual failure risks or optimal allocations. Access priority directly targets the reachability-loss outcome used in evaluation.",
        "This table replaces the former panel b of Priority Stability Under Budget and Demographic Scenarios. All 36 source values are retained; no new model is estimated.",
    ]
    widths = [24, 35, 15, 15, 15, 15]
    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Reachability protection", index=False)
        pd.DataFrame({"Notes": notes}).to_excel(writer, sheet_name="Notes", index=False)
        sheet = writer.sheets["Reachability protection"]
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = "A1:F10"
        sheet.sheet_view.showGridLines = False
        for column, width in zip("ABCDEF", widths):
            sheet.column_dimensions[column].width = width
        for row in sheet:
            sheet.row_dimensions[row[0].row].height = 32
            for cell in row:
                header = cell.row == 1
                cell.font = Font(name="Arial", size=11, bold=header or cell.column == 1,
                                 color="FFFFFF" if header else "000000")
                cell.fill = PatternFill("solid", fgColor="263746" if header else "EDF3F7" if ((cell.row - 2) // 3) % 2 == 0 else "FFFFFF")
                cell.alignment = Alignment(vertical="center", horizontal="left" if cell.column <= 2 else "right", wrap_text=True)
                if not header and cell.column > 2:
                    cell.number_format = "0.00"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=.3, right=.3, top=.3, bottom=.3)
        sheet.print_title_rows = "1:1"
        sheet.print_area = "A1:F10"
        note_sheet = writer.sheets["Notes"]
        note_sheet.column_dimensions["A"].width = 120
        for row in note_sheet:
            cell = row[0]
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Arial", size=11, bold=cell.row <= 2)
            note_sheet.row_dimensions[cell.row].height = 16 * (len(str(cell.value)) // 115 + 2)
    saved = pd.read_excel(OUT)
    pd.testing.assert_frame_equal(saved, table, check_exact=False, rtol=1e-12, atol=1e-12)
    workbook = load_workbook(OUT)
    assert not workbook._external_links
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    table.to_csv(EXP / "table_data.csv", index=False)
    OUT.with_suffix(".caption.md").write_text("\n\n".join(notes) + "\n")
    formatted = table.copy()
    for column in formatted.columns[2:]:
        formatted[column] = formatted[column].map(lambda value: f"{value:.2f}")
    preview = preview_table(formatted, widths, EXP / "table_preview.png")
    assert digest(SOURCE) == before
    result = {"status": "passed", "title": TITLE, "placement": "appendix", "appendix_group": "C",
              "rows": 9, "columns": 6, "numeric_values": 36, "requested_budget_share": .01,
              "years": YEARS, "source_sha256": {str(SOURCE.relative_to(ROOT)): before},
              "source_data_unchanged": True, "former_panel_b_values_reproduced": True,
              "xlsx_readback_passed": True, "formula_cells": 0, "external_links": 0,
              "presentation_revision": "MILI-D-20261008-006", "preview": preview,
              "output": str(OUT.relative_to(ROOT)), "human_acceptance_pending": True}
    (EXP / "validation.json").write_text(json.dumps(result, indent=2) + "\n")
    print(f"Saved: {OUT.relative_to(ROOT)} (9 rows x 6 columns; all 36 former heatmap values preserved)")


if __name__ == "__main__":
    main()

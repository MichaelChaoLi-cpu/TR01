#!/usr/bin/env python3
"""Allocation Rules and Shared Scenario Assumptions.

Plan: ten components, five columns; document adopted assumptions without adding
an analysis. Framework: AnaSOP Sections 5-7, bounded conditional bridge-anchor
comparison, shared cost/action/events, fixed-prefix rules and original area groups.
Run locally: .venv/bin/python src/analyses/table_allocation_rules_and_shared_scenario_assumptions.py
"""
from __future__ import annotations

import hashlib
import json
import math
import os
from pathlib import Path

# Use a writable project cache when running in a restricted local workspace.
os.environ.setdefault("MPLCONFIGDIR", str(Path(__file__).resolve().parents[2] / "data/exp/.matplotlib-cache"))

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[2]
TITLE = "Allocation Rules and Shared Scenario Assumptions"
SLUG = "allocation_rules_and_shared_scenario_assumptions"
EXP = ROOT / "data/exp/figure-table-generation" / SLUG
OUT = ROOT / "data/results/tables" / f"Table_{SLUG}.xlsx"


def digest(path):
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def preview_table(table, widths, preview_path=None):
    """Render the same workbook cells for review, with measured text wrapping."""
    font_size = 9
    fractions = np.asarray(widths) / sum(widths)
    fig = plt.figure(figsize=(16, 10), dpi=100)
    renderer = fig.canvas.get_renderer()
    wrapped = []
    line_counts = []
    for row_index, values in enumerate([list(table.columns)] + table.values.tolist()):
        row, counts = [], []
        for col, value in enumerate(values):
            font = FontProperties(size=font_size, weight="bold" if row_index == 0 or col == 0 else "normal")
            available = 16 * 100 * 0.96 * fractions[col] * 0.88
            lines, line = [], ""
            for word in str(value).split():
                candidate = f"{line} {word}".strip()
                if line and renderer.get_text_width_height_descent(candidate, font, False)[0] > available:
                    lines.append(line)
                    line = word
                else:
                    line = candidate
            lines.append(line)
            row.append("\n".join(lines))
            counts.append(len(lines))
        wrapped.append(row)
        line_counts.append(max(counts))
    row_points = np.asarray(line_counts) * font_size * 1.5 + 20
    fig.set_size_inches(16, (row_points.sum() + 24) / 72)
    ax = fig.add_axes([0.02, 0.015, 0.96, 0.97])
    ax.set_axis_off()
    artist = ax.table(cellText=wrapped, colWidths=fractions, cellLoc="left", bbox=[0, 0, 1, 1])
    artist.auto_set_font_size(False)
    artist.set_fontsize(font_size)
    for (row, col), cell in artist.get_celld().items():
        cell.set_height(row_points[row] / row_points.sum())
        cell.set_facecolor("#263746" if row == 0 else "#EDF3F7" if row % 2 else "white")
        cell.set_edgecolor("white")
        cell.set_linewidth(0.8)
        cell.get_text().set_color("white" if row == 0 else "black")
        cell.get_text().set_fontweight("bold" if row == 0 or col == 0 else "normal")
        cell.PAD = 0.06
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    overflow = []
    for key, cell in artist.get_celld().items():
        text_box = cell.get_text().get_window_extent(renderer)
        cell_box = cell.get_window_extent(renderer)
        if text_box.x0 < cell_box.x0 or text_box.x1 > cell_box.x1 or text_box.y0 < cell_box.y0 or text_box.y1 > cell_box.y1:
            overflow.append(key)
    assert not overflow, f"Preview cell overflow: {overflow}"
    preview = Path(preview_path) if preview_path is not None else EXP / "table_preview.png"
    fig.savefig(preview, dpi=600, facecolor="white")
    size = fig.get_size_inches().tolist()
    plt.close(fig)
    return {"path": str(preview.relative_to(ROOT)), "dpi": 600, "size_inches": size, "text_overflow_cells": overflow}


def main():
    unit_path = ROOT / "data/processed/common-budget/conditional_maintenance_units_preprocessed.parquet"
    curve_path = ROOT / "data/processed/common-budget/conditional_protection_curves_preprocessed.parquet"
    construction_path = ROOT / "data/exp/data-preprocessing/common-budget/construction_validation.json"
    allocation_path = ROOT / "data/exp/common-budget-allocation/validation.json"
    plan_path = ROOT / "docs/AnaSOP.md"
    paths = [unit_path, curve_path, construction_path, allocation_path, plan_path]
    hashes = {str(p.relative_to(ROOT)): digest(p) for p in paths}
    units = pd.read_parquet(unit_path)
    curves = pd.read_parquet(curve_path)
    construction = json.loads(construction_path.read_text())
    allocation = json.loads(allocation_path.read_text())
    plan = plan_path.read_text()
    assert TITLE in plan and "| State shared intervention assumptions |" in plan
    assert all(f"## {i}." in plan for i in range(4, 9))
    assert construction["status"] == allocation["status"] == "passed"
    n = len(units)
    traffic_n = units["Traffic Observation Key"].nunique()
    assert n == construction["common_units"] == allocation["common_units"] == 46688
    assert units["Graph Source Edge ID"].is_unique and units["Bridge Record ID"].is_unique
    assert units["Source Prefecture"].nunique() == 47
    assert traffic_n == construction["common_traffic_units"] == 15815
    assert units["Policy Common Sample Eligible"].all()
    # This source-preserving descriptor is stored as text, not a Boolean.
    assert units["Physical Asset Match Accepted"].astype(str).str.lower().eq("false").all()
    cost = units["Relative Intervention Cost"]
    assert cost.gt(0).all()
    np.testing.assert_allclose(cost, units["Bridge Length"] * units["Bridge Width"], rtol=1e-12)
    np.testing.assert_allclose(cost.sum(), allocation["relative_total_cost"], rtol=1e-12)
    assert units["Maintenance Intervention Effect"].eq(1).all()
    np.testing.assert_allclose(units["Conditional Event Weight"], 1 / n, rtol=1e-12)
    assert units["Bridge Inspection Category"].isin([1, 2, 3, 4]).all()
    assert units["Observed Traffic Volume"].notna().all()
    for _, frame in curves.groupby("Maintenance Priority"):
        np.testing.assert_allclose(frame["Budget Share"], np.arange(101) / 100, atol=1e-14)
    assert len(curves) == 303 and curves["Maintenance Priority"].nunique() == 3
    age_cutoff = construction["ageing_cutoff"] * 100
    distance_cutoff = construction["baseline_distance_cutoff_metres"] / 1000
    rows = [
        ["Common sample",
         f"{n:,} units; 47 prefectures.",
         "Observed inputs; assumed linkage.",
         "Same units, origins and destinations.",
         "Restricted common sample."],
        ["Association and loss event",
         "One bridge–road link; single-road closure.",
         "Assumed link; computed loss.",
         "Fixed entries; alternatives retained.",
         "Physical match unverified."],
        ["Relative cost",
         "Bridge length × width.",
         "Observed dimensions; assumed unit cost.",
         "Same unit costs.",
         "Relative, not monetary."],
        ["Effect and event weights",
         "Full prevention; equal single-event weights.",
         "Assumed effect and weights.",
         "Same effect and weights.",
         "Conditional; no annual risk."],
        ["Budget and selection",
         "0–100%, 1-point steps; longest affordable prefix.",
         "Adopted algorithm.",
         "Same requested budget.",
         "Indivisible units; unused budget allowed."],
        ["Traffic priority",
         "Traffic / cost (descending).",
         "Measured traffic; adopted rule.",
         "Common cost and tie-breaker.",
         "Shared traffic observations."],
        ["Condition priority",
         "IV first, followed by III, II and I; then lower cost.",
         "Observed ordinal categories; adopted rule.",
         "Common cost and tie-breaker.",
         "Mixed inspection years."],
        ["Access priority",
         "Lost-path population / cost; then added-distance burden / cost (descending).",
         "Computed loss; adopted rule.",
         "Common cost and tie-breaker.",
         "Outcome-aligned; equity not presumed."],
        ["Area groups",
         "Upper quartile of ageing or baseline medical distance.",
         "Observed baseline; fixed cutoffs.",
         "Same origin-based groups.",
         "Overlapping area groups."],
        ["Demographic extension",
         "2020, 2030, 2040 and 2050 matched-population sensitivity.",
         "Official mesh changes; census-anchored weights.",
         "Fixed system, units and baseline area groups.",
         "Separate sensitivity; no future elderly locations."],
    ]
    table = pd.DataFrame(rows, columns=["Component", "Definition", "Observation/assumption basis", "Commonality across rules", "Interpretation limit"])
    assert table.shape == (10, 5) and table.notna().all().all()
    detailed_notes = [
        TITLE,
        "This Appendix C table documents the adopted framework; it introduces no new estimation, intervention assumption or demographic scenario. Observed source inputs are distinguished from conditional associations, computed outcomes and adopted evaluation rules.",
        "Traffic priority sorts observed traffic per relative cost descending. Condition priority sorts categories IV, III, II and I, then cost ascending. Access priority sorts lost-path population per cost descending, then additional distance burden per cost descending. Exact final ties use ascending stable source-road identity.",
        "Bridge association uses the adopted nearest eligible anchor within 50 m, a single candidate within 5 m of the nearest distance, and the confirmed quality/conflict exclusions. Shared nearest-road cases are withheld globally. No complete physical bridge-to-road intervention match has been verified.",
        "For each single-source-road event, lost-path population is the population without a remaining finite medical path. Additional distance burden is population times added distance, in person-kilometres, only for origins with a remaining finite path. Disconnected origins receive no arbitrary finite distance penalty. Required closed-road entries are removed without reattachment; parallel alternatives remain.",
        "The baseline retains 175,612 mesh origins and 125,874,788 residents. Population and registered medical facilities refer to 2020, candidate roads to 2024 and observed traffic to 2021. Inspection dates vary across retained assets. Candidate access does not establish current medical service availability or legal/time-based routing.",
        f"The relative-cost sum is {float(cost.sum()):,.2f} square-metre equivalents under a common unit price. Budget is a fraction of this sum, not a monetary allocation. Indivisible fixed-prefix selection can leave budget unused; it is not an optimization solution.",
        "Aggregate or group protection equals selected-event loss divided by total loss across the same mutually exclusive event domain, separately for reachability and additional finite-route distance. Uniform event weights cancel in these fractions. Conditional avoided loss per event divides selected-event loss by the number of events. Undefined zero-denominator fractions remain missing.",
        f"Area cutoffs are the unweighted 75th percentiles on 145,771 original selected official publication supports, with ties retained: {age_cutoff:.9f}% ageing among known ages and {distance_cutoff:.9f} km population-weighted baseline medical distance. Meshes inherit original support membership; older residents are not redistributed. Higher-ageing areas contain {construction['higher_ageing_population']:,} selected residents and distance-disadvantaged areas {construction['distance_disadvantaged_population']:,}; overlap is {construction['overlap_population']:,}. Each group and its complement reconcile separately to event totals.",
        "Access priority is deliberately aligned with evaluated reachability loss. Its performance does not independently establish an optimal allocation, equity, empirical maintenance effectiveness, or superiority on unmeasured condition-risk objectives.",
        "Baseline reporting budgets are 1%, 5%, 10% and 20% of total relative cost. The separate adopted bridge sensitivity uses 175,485 forecast-matched origins and 125,874,510 census residents at matched 2020, excluding 127 unmatched origins and 278 residents only from the demographic comparison. Original census population is reweighted by official mesh projected-to-2020 population ratios for 2030, 2040 and 2050. Units, costs, network, destinations and original 2020 area groups stay fixed; access priorities are reranked. This does not infer future elderly-resident locations or pavement sensitivity.",
    ]
    notes = [
        TITLE,
        "The same units, origins, medical destinations, costs, maintenance effect and event weights apply to all three rules. Only priority order changes. Exact final ties use stable source-road identity.",
        f"Traffic is measured over daytime 12 hours; {traffic_n:,} observation units are shared across {n:,} roads. Condition categories are ordinal. Access priority ranks medical reachability loss first, then added distance among still-reachable origins.",
        "Bridge-road linkage, area-proportional cost, complete event prevention and equal mutually exclusive event weights are assumptions. Protection is conditional, not annual risk, unique beneficiaries or a proved optimum. Unaffordable units are neither skipped nor split; report actual and unused spending.",
        f"Area groups use original-support upper-quartile cutoffs ({age_cutoff:.3f}% ageing; {distance_cutoff:.3f} km baseline medical distance), retaining ties. Groups overlap and describe all residents in those areas. No older residents are redistributed. The separate bridge demographic sensitivity uses official spatial population-change factors on matched origins and fixed original area groups.",
    ]
    EXP.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    widths = [24, 47, 34, 31, 32]
    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Rules and assumptions", index=False)
        pd.DataFrame({"Notes": notes}).to_excel(writer, sheet_name="Notes", index=False)
        pd.DataFrame({"Detailed methods": detailed_notes}).to_excel(writer, sheet_name="Detailed methods", index=False)
        sheet = writer.sheets["Rules and assumptions"]
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = "A1:E11"
        sheet.sheet_view.showGridLines = False
        for column, width in zip("ABCDE", widths):
            sheet.column_dimensions[column].width = width
        for row in sheet:
            header = row[0].row == 1
            line_count = max(sum(max(1, math.ceil(len(line) / (width - 4))) for line in str(cell.value).splitlines()) for cell, width in zip(row, widths))
            sheet.row_dimensions[row[0].row].height = max(35, line_count * 16 + 18)
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="263746" if header else "EDF3F7" if cell.row % 2 == 0 else "FFFFFF")
                cell.font = Font(name="Arial", size=11, bold=header or cell.column == 1, color="FFFFFF" if header else "000000")
                cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
        sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.3, bottom=0.3)
        sheet.print_title_rows = "1:1"
        sheet.print_area = "A1:E11"
        sheet.print_options.horizontalCentered = True
        for name in ["Notes", "Detailed methods"]:
            note_sheet = writer.sheets[name]
            note_sheet.column_dimensions["A"].width = 125
            note_sheet.sheet_view.showGridLines = False
            for row in note_sheet:
                cell = row[0]
                cell.font = Font(name="Arial", size=11, bold=cell.row <= 2, color="000000")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                note_sheet.row_dimensions[cell.row].height = max(24, math.ceil(len(str(cell.value)) / 120) * 16 + 16)
    saved = pd.read_excel(OUT, sheet_name="Rules and assumptions")
    pd.testing.assert_frame_equal(saved, table, check_exact=True)
    book = load_workbook(OUT, data_only=False)
    assert book["Rules and assumptions"].max_row == 11 and book["Rules and assumptions"].max_column == 5
    assert [row[0].value for row in book["Notes"].iter_rows(min_row=2)] == notes
    assert [row[0].value for row in book["Detailed methods"].iter_rows(min_row=2)] == detailed_notes
    assert not book._external_links
    assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)
    table.to_csv(EXP / "table_data.csv", index=False)
    OUT.with_suffix(".caption.md").write_text("\n\n".join(notes) + "\n")
    preview = preview_table(table, widths)
    assert all(digest(ROOT / name) == expected for name, expected in hashes.items())
    validation = {
        "title": TITLE, "status": "generated_awaiting_human_review", "rows": 10, "columns": 5,
        "placement": "appendix", "appendix_group": "C", "framework_authorization": "MILI-D-20261007-026",
        "framework": "AnaSOP Sections 5-7: shared conditional association, cost, action, weights, rankings and fixed-prefix budgets",
        "source_hashes_verified_after_export": hashes, "input_data_unchanged": True,
        "common_units": n, "shared_traffic_measurement_units": traffic_n,
        "cost_formula_effect_weights_and_budget_grid_checks": "passed", "xlsx_exact_readback": "passed",
        "observations_separated_from_assumptions": True, "demographic_scenario": "adopted_matched_population_sensitivity", "demographic_authorization": "MILI-D-20261007-036",
        "new_estimation_or_assumptions": False, "formula_cells": 0, "external_links": 0,
        "output": str(OUT.relative_to(ROOT)),
        "preview": preview,
        "presentation_revision": "MILI-D-20261007-030",
        "main_table_word_count": sum(len(str(cell).split()) for row in table.values for cell in row),
        "detailed_methods_preserved": True,
    }
    (EXP / "validation.json").write_text(json.dumps(validation, indent=2) + "\n")
    print(f"Saved: {OUT.relative_to(ROOT)}")
    print("Validated: 10 × 5 cells, exact Excel read-back, shared inputs and unchanged source hashes.")


if __name__ == "__main__":
    main()

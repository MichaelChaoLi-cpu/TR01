"""Pavement Maintenance Protection and Resource Accounting.

AnaSOP joint finite-cost outcomes, fixed original area groups and complete
prefix resources, 25% baseline increment. Local pandas/openpyxl reproduction.
"""
from pathlib import Path
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font,PatternFill

ROOT=Path(__file__).resolve().parents[2]
SLUG='pavement_maintenance_protection_and_resource_accounting'
EXP=ROOT/'data/exp/figure-table-generation'/SLUG
OUT=ROOT/'data/results/tables'/f'Table_{SLUG}.xlsx'


def main():
    EXP.mkdir(parents=True,exist_ok=True);OUT.parent.mkdir(parents=True,exist_ok=True)
    df=pd.read_parquet(ROOT/'data/processed/pavement-policy/pavement_group_outcomes_preprocessed.parquet')
    df=df[df['Pavement Impedance Increment'].eq(.25)]
    rows=[]
    for share in [.01,.05,.10,.20]:
        for rule in ['Traffic','Condition','Access']:
            group=df[df['Pavement Budget Share'].eq(share)&df['Pavement Maintenance Priority'].eq(rule)].set_index('Pavement Population Group')
            all_res=group.loc['All Residents']
            rows.append({'Requested budget (%)':share*100,'Priority':rule,
                         'Actual spending (%)':all_res['Pavement Actual Spending Share']*100,
                         'Unused budget (%)':all_res['Pavement Unused Budget Share']*100,
                         'Selected roads':int(all_res['Pavement Selected Road Count']),
                         'Avoided burden (resident reference-km)':all_res['Pavement Avoided Access Burden'],
                         'All residents protected (%)':all_res['Pavement Access Protection Share']*100,
                         'Higher-ageing areas protected (%)':group.loc['Higher Ageing Area','Pavement Access Protection Share']*100,
                         'Distance-disadvantaged areas protected (%)':group.loc['Baseline Distance Disadvantaged Area','Pavement Access Protection Share']*100})
    table=pd.DataFrame(rows);assert table.shape==(12,9)
    notes=[('Reference','25% unmaintained-road impedance increment; fixed original population and area masks.'),
           ('Budget','Share of the summed length-relative cost of 70,049 conditional pavement road units.'),
           ('Selection','Longest affordable complete priority prefix, without splitting or skipping units.'),
           ('Protection','Avoided population-weighted excess medical-access cost divided by the same group all-unmaintained reference.'),
           ('Units','Avoided burden is in resident reference-kilometre equivalents, not observed journeys or travel time.'),
           ('Area groups','Original official support labels; overlapping groups are not additive and do not locate individual older residents.'),
           ('Precision','Workbook stores full validated values; displayed rounding does not alter calculations.')]
    with pd.ExcelWriter(OUT,engine='openpyxl') as writer:
        table.to_excel(writer,index=False,sheet_name='Results')
        pd.DataFrame(notes,columns=['Item','Definition']).to_excel(writer,index=False,sheet_name='Notes')
    book=load_workbook(OUT);ws=book['Results'];ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
    widths=[15,13,15,15,14,24,20,22,25]
    for i,width in enumerate(widths,1):ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=width
    ws.row_dimensions[1].height=48
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='35566C');cell.alignment=Alignment(wrap_text=True,vertical='center')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(vertical='center',horizontal='left' if cell.column==2 else 'right')
            if cell.column in [3,4]:cell.number_format='0.0000'
            elif cell.column in [1,7,8,9]:cell.number_format='0.00'
            elif cell.column in [5,6]:cell.number_format='#,##0' if cell.column==5 else '#,##0.00'
            if (cell.row-2)//3%2:cell.fill=PatternFill('solid',fgColor='EEF3F6')
        ws.row_dimensions[row[0].row].height=21
    book['Notes'].column_dimensions['A'].width=18;book['Notes'].column_dimensions['B'].width=105
    for row in book['Notes'].iter_rows():
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    book.save(OUT)
    check=pd.read_excel(OUT,sheet_name='Results')
    assert list(check.columns)==list(table.columns) and check.shape==table.shape
    for col in table:
        if col=='Priority':assert check[col].tolist()==table[col].tolist()
        else:np.testing.assert_allclose(check[col],table[col],rtol=1e-12,atol=1e-10)
    table.to_csv(EXP/'table_values.csv',index=False)
    # Optional local preview; workbook and source accounting remain authoritative.
    import os
    os.environ.setdefault('MPLCONFIGDIR','/private/tmp/tr01-matplotlib')
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    headers=['Budget\n(%)','Priority','Actual\n(%)','Unused\n(%)','Selected\nroads',
             'Avoided burden\n(resident reference-km)','All residents\nprotected (%)',
             'Higher-ageing\nprotected (%)','Disadvantaged\nprotected (%)']
    preview_rows=[]
    for row in table.to_dict('records'):
        cells=[]
        for i,value in enumerate(row.values()):
            cells.append(str(value) if i==1 else f'{value:,.0f}' if i==4 else
                         f'{value:,.4f}' if i in [2,3] else f'{value:,.2f}')
        preview_rows.append(cells)
    fig,ax=plt.subplots(figsize=(13,4.8));ax.axis('off')
    preview=ax.table(cellText=preview_rows,colLabels=headers,cellLoc='right',colLoc='center',
                     bbox=[0,0,1,1],colWidths=[.055,.075,.07,.07,.07,.21,.14,.15,.16])
    preview.auto_set_font_size(False);preview.set_fontsize(8)
    for (r,c),cell in preview.get_celld().items():
        cell.set_edgecolor('#CAD4DA');cell.set_linewidth(.4)
        if r==0:
            cell.set_facecolor('#35566C');cell.set_text_props(color='white',weight='bold')
            cell.set_height(cell.get_height()*1.4)
        elif (r-1)//3%2:cell.set_facecolor('#EEF3F6')
    fig.subplots_adjust(left=.01,right=.99,top=.99,bottom=.01)
    fig.savefig(EXP/'table_preview.png',dpi=150);plt.close(fig)
    OUT.with_suffix('.caption.md').write_text('''# Pavement Maintenance Protection and Resource Accounting

Three fixed maintenance rules are evaluated at requested budgets of 1, 5, 10 and 20 percent of total pavement-domain length-relative cost, under the 25 percent baseline impedance increment. Actual and unused spending refer to the same total cost denominator and reflect indivisible complete-prefix selection. Protection fractions compare avoided excess generalized medical-access burden with each group's all-unmaintained reference; group denominators differ and groups overlap. Avoided burden uses resident reference-kilometre equivalents. Results concern the conditional road sample and adopted finite-cost assumptions, not observed expenditure, actual journeys, individual older residents or causal repair effects. Full numeric precision is retained in the workbook, with notes on the second sheet.
''')
    (EXP/'validation.json').write_text(json.dumps({'status':'passed','rows':12,'columns':9,
                'workbook_readback_exact_within_float_tolerance':True,'budget_rule_grid_complete':True,
                'framework_followed':True,'human_acceptance_pending':True},indent=2)+'\n')
    print(f'Saved {OUT.relative_to(ROOT)} (12 rows x 9 columns)')


if __name__=='__main__':main()

#!/usr/bin/env python3
"""Baseline Analysis Sample Selection.

Plan: seven-row accounting of the original meshes and overlapping exclusions.
Framework: AnaSOP Sections 5-7, confirmed selection masks and baseline population
shares; preserve publication supports and distinguish overlap from the union.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
TITLE = "Baseline Analysis Sample Selection"
SLUG = "baseline_analysis_sample_selection"
BASE = ROOT / "data/processed/medical-baseline"
SAMPLE = ROOT / "data/processed/medical-analysis-sample"
EXP = ROOT / "data/exp/figure-table-generation" / SLUG
OUT = ROOT / "data/results/tables" / f"Table_{SLUG}.xlsx"
KEY = "Publication Support Code"
AGE = "Population Aged 65+ Share Among Known Ages"
DISTANCE = "Candidate Medical Network Distance"
POP = "Local Total Population"


def digest(path):
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def main():
    paths = [BASE / "medical_mesh_baseline.parquet",
             BASE / "medical_age_support_baseline.parquet",
             SAMPLE / "medical_mesh_analysis_preprocessed.parquet",
             SAMPLE / "medical_age_support_analysis_preprocessed.parquet",
             ROOT / "data/exp/data-preprocessing/medical-analysis-sample/sample_accounting.csv"]
    hashes = {str(p.relative_to(ROOT)): digest(p) for p in paths}
    meshes, supports, selected_meshes, selected_supports = [pd.read_parquet(p) for p in paths[:4]]
    assert meshes["Mesh Code"].is_unique and supports[KEY].is_unique
    assert meshes[KEY].isin(supports[KEY]).all()
    assert meshes[POP].notna().all() and (meshes[POP] >= 0).all()
    age_missing = meshes[KEY].isin(supports.loc[supports[AGE].isna(), KEY])
    distance_missing = meshes[DISTANCE].isna()
    no_path = meshes["Medical Route Status"].eq("eligible_entry_no_candidate_path")
    pending = meshes["Medical Route Status"].eq("entry_pending")
    assert not (no_path & pending).any()
    np.testing.assert_array_equal(distance_missing, no_path | pending)
    overlap = age_missing & distance_missing
    excluded = age_missing | distance_missing
    pd.testing.assert_frame_equal(selected_meshes, meshes.loc[~excluded].reset_index(drop=True))
    support_excluded = supports[AGE].isna() | supports["Population-Weighted Candidate Medical Network Distance"].isna()
    pd.testing.assert_frame_equal(selected_supports, supports.loc[~support_excluded].reset_index(drop=True))
    assert set(selected_meshes[KEY]) == set(selected_supports[KEY])
    population = int(meshes[POP].sum())
    masks = {
        "All baseline meshes": np.ones(len(meshes), dtype=bool),
        "Undefined ageing share": age_missing,
        "Eligible road entry without baseline medical path": no_path,
        "Pending baseline road entry": pending,
        "Overlap of age and distance exclusions": overlap,
        "Excluded union": excluded,
        "Retained analysis sample": ~excluded,
    }
    rows = []
    for category, mask in masks.items():
        residents = int(meshes.loc[mask, POP].sum())
        rows.append({"Category": category, "Unit": "mesh", "Mesh count": int(np.sum(mask)),
                     "Resident count": residents, "Share of baseline population": residents / population})
    table = pd.DataFrame(rows)
    indexed = table.set_index("Category")
    for col in ["Mesh count", "Resident count"]:
        components = indexed.loc[["Undefined ageing share", "Eligible road entry without baseline medical path", "Pending baseline road entry"], col].sum()
        assert indexed.loc["Excluded union", col] == components - indexed.loc["Overlap of age and distance exclusions", col]
        assert indexed.loc["Excluded union", col] + indexed.loc["Retained analysis sample", col] == indexed.loc["All baseline meshes", col]
    prior = pd.read_csv(paths[-1]).rename(columns={"Record Count": "Mesh count", "Resident Count": "Resident count", "Share of Baseline Population": "Share of baseline population"})
    pd.testing.assert_frame_equal(table, prior, check_exact=False, rtol=1e-12, atol=1e-15)
    assert len(meshes) == 176964 and population == 126146099
    assert len(selected_meshes) == 175612 and len(selected_supports) == 145771
    assert int(selected_meshes[POP].sum()) == 125874788
    notes = [
        TITLE,
        "Each count refers to observed baseline meshes; resident counts sum Local Total Population. All population shares use the original baseline population of 126,146,099 residents as the denominator.",
        "Exclusion components overlap. Subtract the age/distance overlap once when calculating the excluded union. The overlap is a subset, not an additional exclusion category.",
        "Undefined ageing share excludes all members of the affected official publication support. Undefined medical distance includes eligible entries without a candidate path and pending entries; pending entry is not verified physical disconnection.",
        "The retained mesh population is 125,874,788. The selected 145,771 age publication supports retain their original age and full-support population counts; full-support population is not the selected mesh population. No older residents are redistributed to member meshes.",
        "Candidate medical access uses registered hospitals and general clinics, not verified current service availability. Population and facilities refer to 2020 and the candidate road configuration to 2024.",
        "Missing baseline distances are not replaced by zero. Scenario-induced disconnection of an originally eligible origin would be a policy outcome, not a new baseline exclusion. This table describes sample coverage, not maintenance benefits.",
        "Sources: e-Stat census population and official age publication supports; MLIT National Land Numerical Information medical facilities and candidate road centre lines.",
    ]
    EXP.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Sample selection", index=False)
        pd.DataFrame({"Notes": notes}).to_excel(writer, sheet_name="Notes", index=False)
        sheet = writer.sheets["Sample selection"]
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = "A1:E8"
        for col, width in zip("ABCDE", [58, 12, 18, 20, 33]):
            sheet.column_dimensions[col].width = width
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="263746")
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2, max_row=8):
            bold = row[0].value in ["All baseline meshes", "Excluded union", "Retained analysis sample"]
            colour = "FFF2D9" if row[0].value == "Excluded union" else "EDF3F7" if row[0].row % 2 == 0 else "FFFFFF"
            sheet.row_dimensions[row[0].row].height = 34
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=colour)
                cell.font = Font(name="Arial", size=11, color="000000", bold=bold)
                cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="right" if cell.column >= 3 else "left")
            row[2].number_format = row[3].number_format = "#,##0"
            row[4].number_format = "0.000000%"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:1"
        sheet.print_area = "A1:E8"
        note_sheet = writer.sheets["Notes"]
        note_sheet.column_dimensions["A"].width = 125
        for row in note_sheet:
            row[0].font = Font(name="Arial", size=11, color="000000", bold=row[0].row <= 2)
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
            note_sheet.row_dimensions[row[0].row].height = 58 if row[0].row > 2 else 24
    saved = pd.read_excel(OUT, sheet_name="Sample selection")
    pd.testing.assert_frame_equal(saved, table, check_exact=False, rtol=1e-12, atol=1e-15)
    book = load_workbook(OUT, data_only=True)
    assert book["Sample selection"].max_row == 8 and book["Sample selection"].max_column == 5
    assert all(book["Sample selection"][f"E{row}"].number_format == "0.000000%" for row in range(2, 9))
    assert [book["Notes"][f"A{row}"].value for row in range(2, 10)] == notes
    table.to_csv(EXP / "table_data.csv", index=False)
    OUT.with_suffix(".caption.md").write_text("\n\n".join(notes) + "\n")
    for name, expected in hashes.items():
        assert digest(ROOT / name) == expected
    receipt = {
        "title": TITLE, "status": "generated_awaiting_human_review", "rows": 7, "columns": 5,
        "framework": "AnaSOP Sections 5-7: original baseline, overlapping masks, exclusion union and population shares",
        "sample_authorization": "MILI-D-20261006-019", "inventory_authorization": "MILI-D-20261006-020",
        "population_share_denominator": population, "source_hashes_verified_after_export": hashes,
        "baseline_meshes": len(meshes), "retained_meshes": len(selected_meshes), "retained_age_supports": len(selected_supports),
        "excluded_union_meshes": int(excluded.sum()), "excluded_union_residents": int(meshes.loc[excluded, POP].sum()),
        "overlap_meshes": int(overlap.sum()), "overlap_residents": int(meshes.loc[overlap, POP].sum()),
        "independent_masks_and_prior_accounting_check": "passed", "inclusion_exclusion_identity": "passed",
        "exact_selected_mesh_and_support_subsets": "passed", "xlsx_readback": "passed",
        "input_data_unchanged": True, "age_redistribution": False, "output": str(OUT.relative_to(ROOT)),
        "population_share_display": "Six decimal percentage places; stored fractions retain full precision",
    }
    (EXP / "validation.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(f"Saved: {OUT.relative_to(ROOT)}")
    print(table.to_string(index=False))
    print("Validated: masks, overlapping union, population conservation, original selected subsets and Excel read-back.")


if __name__ == "__main__":
    main()

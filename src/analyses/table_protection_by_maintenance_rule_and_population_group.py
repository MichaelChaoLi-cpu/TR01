#!/usr/bin/env python3
"""Protection by Maintenance Rule and Population Group.

Framework: AnaSOP Sections 5-7, three fixed priority lists, common requested
relative budgets, separate conditional protection outcomes and origin groups.
Budget rows require an explicit human decision; this script has no default grid.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / "data/exp/.matplotlib-cache"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

TITLE = "Protection by Maintenance Rule and Population Group"
SLUG = "protection_by_maintenance_rule_and_population_group"
EXP = ROOT / "data/exp/figure-table-generation" / SLUG
OUT = ROOT / "data/results/tables" / f"Table_{SLUG}.xlsx"
RULES = ["Traffic priority", "Condition priority", "Access priority"]
PROTECTION = {
    "Distance burden protection": "Distance Burden Protection Share",
    "Reachability protection": "Reachability Protection Share",
    "Higher-ageing area protection": "Higher Ageing Reachability Protection Share",
    "Distance-disadvantaged area protection": "Distance Disadvantaged Reachability Protection Share",
}


def digest(path):
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def construct_table(curves, budgets, total_cost):
    rows = []
    for percent in budgets:
        share = percent / 100
        subset = curves.loc[np.isclose(curves["Budget Share"], share, rtol=0, atol=1e-12)]
        assert len(subset) == 3 and set(subset["Maintenance Priority"]) == set(RULES)
        for rule in RULES:
            source = subset.loc[subset["Maintenance Priority"].eq(rule)].iloc[0]
            spent = float(source["Actual Relative Spending"] / total_cost)
            unused = float(source["Unused Relative Budget"] / total_cost)
            np.testing.assert_allclose(spent + unused, share, rtol=1e-10, atol=1e-12)
            assert spent <= share + 1e-12 and unused >= 0
            row = {"Requested budget": share, "Maintenance rule": rule,
                   "Actual spending": spent, "Unused budget": unused}
            for label, column in PROTECTION.items():
                value = float(source[column])
                assert np.isfinite(value) and 0 <= value <= 1 + 1e-12
                row[label] = value
            rows.append(row)
    return pd.DataFrame(rows)


def render_preview(table):
    headers = ["Requested\nbudget (%)", "Maintenance\nrule", "Actual\nspending (%)", "Unused\nbudget (%)",
               "Distance burden\nprotection (%)", "Reachability\nprotection (%)",
               "Higher-ageing\narea protection (%)", "Distance-disadvantaged\narea protection (%)"]
    values = []
    for row in table.itertuples(index=False, name=None):
        values.append([f"{100 * row[0]:g}", row[1].replace(" priority", ""),
                       f"{100 * row[2]:.4f}", f"{100 * row[3]:.4f}",
                       *[f"{100 * value:.2f}" for value in row[4:]]])
    fig = plt.figure(figsize=(12, 0.39 * len(table) + 0.9))
    ax = fig.add_axes([0.02, 0.025, 0.96, 0.95])
    ax.set_axis_off()
    artist = ax.table(cellText=values, colLabels=headers,
                      colWidths=[0.10, 0.12, 0.10, 0.10, 0.13, 0.12, 0.16, 0.17],
                      cellLoc="center", bbox=[0, 0, 1, 1])
    artist.auto_set_font_size(False)
    artist.set_fontsize(9)
    for (row, col), cell in artist.get_celld().items():
        header = row == 0
        group = (row - 1) // 3
        cell.set_facecolor("#263746" if header else "#EDF3F7" if group % 2 == 0 else "white")
        cell.set_edgecolor("white")
        cell.set_linewidth(0.8)
        cell.get_text().set_color("white" if header else "black")
        cell.get_text().set_fontweight("bold" if header else "normal")
        cell.set_height(1.65 / (len(table) + 1.65) if header else 1 / (len(table) + 1.65))
        cell.PAD = 0.025
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    overflow = []
    for key, cell in artist.get_celld().items():
        text_box = cell.get_text().get_window_extent(renderer)
        box = cell.get_window_extent(renderer)
        if text_box.x0 < box.x0 or text_box.x1 > box.x1 or text_box.y0 < box.y0 or text_box.y1 > box.y1:
            overflow.append(key)
    assert not overflow, f"Preview text overflow: {overflow}"
    fig.savefig(EXP / "table_preview.png", dpi=600, facecolor="white")
    plt.close(fig)
    return {"dpi": 600, "text_overflow_cells": overflow}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--budget-percentages", type=float, nargs="+", required=True)
    parser.add_argument("--decision-id", required=True, help="Human decision adopting these reporting rows")
    args = parser.parse_args()
    budgets = args.budget_percentages
    assert budgets == sorted(set(budgets)) and all(0 <= b <= 100 and b == int(b) for b in budgets)
    assert args.decision_id in (ROOT / "docs/framelog/mili.md").read_text()
    plan = (ROOT / "docs/AnaSOP.md").read_text()
    plan_row = next(line for line in plan.splitlines() if line.startswith(f"| {TITLE} |"))
    assert "selection remains pending" not in plan_row, "Confirm reference budgets through figure-table-planning first."
    curve_path = ROOT / "data/processed/common-budget/conditional_protection_curves_preprocessed.parquet"
    allocation_path = ROOT / "data/exp/common-budget-allocation/validation.json"
    paths = [curve_path, allocation_path, ROOT / "docs/AnaSOP.md"]
    hashes = {str(p.relative_to(ROOT)): digest(p) for p in paths}
    curves = pd.read_parquet(curve_path)
    allocation = json.loads(allocation_path.read_text())
    assert allocation["status"] == "passed" and allocation["common_units"] == 46688
    table = construct_table(curves, budgets, allocation["relative_total_cost"])
    assert table.shape == (3 * len(budgets), 8)
    notes = [
        TITLE,
        "Each requested budget compares the same three fixed priority rules on 46,688 conditional bridge-associated road units. Budget and spending percentages use total common-domain relative cost; protection percentages use the respective aggregate or group loss denominator.",
        "The two area columns report reachability protection for all selected residents in higher-ageing and baseline-distance-disadvantaged areas. These overlapping groups are defined on original publication supports; they are not individual older-resident outcomes.",
        "Maintenance fully prevents a selected unit's own event under equally weighted, mutually exclusive single-road events. Cost is bridge length times width. Fixed-prefix selection does not skip or split an unaffordable unit; actual and unused spending can differ among rules.",
        "Protection is conditional on the adopted sample and assumptions. Access priority directly targets evaluated reachability loss; its performance alone does not establish optimality or equity. No annual risk, unique-beneficiary or spatial demographic-change estimate is reported.",
        "Workbook cells retain unrounded fractions. Protection is displayed to two percentage decimals; actual and unused spending to four. Display rounding does not modify source values.",
    ]
    EXP.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Protection", index=False)
        pd.DataFrame({"Notes": notes}).to_excel(writer, sheet_name="Notes", index=False)
        sheet = writer.sheets["Protection"]
        for col, width in zip("ABCDEFGH", [19, 24, 18, 18, 23, 21, 27, 30]):
            sheet.column_dimensions[col].width = width
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = f"A1:H{len(table)+1}"
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            header = row[0].row == 1
            group = (row[0].row - 2) // 3
            sheet.row_dimensions[row[0].row].height = 50 if header else 28
            for cell in row:
                cell.font = Font(name="Arial", size=11, bold=header, color="FFFFFF" if header else "000000")
                cell.fill = PatternFill("solid", fgColor="263746" if header else "EDF3F7" if group % 2 == 0 else "FFFFFF")
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if cell.column == 2 else "center")
                if not header and cell.column != 2:
                    cell.number_format = "0%" if cell.column == 1 else "0.0000%" if cell.column in [3, 4] else "0.00%"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
        sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.3, bottom=0.3)
        sheet.print_title_rows = "1:1"
        sheet.print_area = f"A1:H{len(table)+1}"
        note_sheet = writer.sheets["Notes"]
        note_sheet.column_dimensions["A"].width = 125
        note_sheet.sheet_view.showGridLines = False
        for row in note_sheet:
            row[0].font = Font(name="Arial", size=11, color="000000", bold=row[0].row <= 2)
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
            note_sheet.row_dimensions[row[0].row].height = 66 if row[0].row > 2 else 25
    saved = pd.read_excel(OUT, sheet_name="Protection")
    pd.testing.assert_frame_equal(saved, table, check_exact=False, rtol=1e-14, atol=1e-15)
    book = load_workbook(OUT)
    assert book["Protection"].max_row == len(table)+1 and book["Protection"].max_column == 8
    assert not book._external_links and not any(c.data_type == "f" for sheet in book for row in sheet for c in row)
    assert [row[0].value for row in book["Notes"].iter_rows(min_row=2)] == notes
    table.to_csv(EXP / "table_data.csv", index=False)
    OUT.with_suffix(".caption.md").write_text("\n\n".join(notes)+"\n")
    preview = render_preview(table)
    assert all(digest(ROOT / name) == expected for name, expected in hashes.items())
    receipt = {"title": TITLE, "status": "generated_awaiting_human_review", "rows": len(table), "columns": 8,
               "placement": "main_manuscript", "reference_budget_percentages": budgets,
               "reference_budget_decision": args.decision_id, "framework_authorization": "MILI-D-20261007-026",
               "xlsx_readback": "passed", "budget_resource_identity": "passed",
               "source_hashes_verified_after_export": hashes, "input_data_unchanged": True,
               "new_estimation_or_assumptions": False, "preview": preview, "output": str(OUT.relative_to(ROOT))}
    (EXP / "validation.json").write_text(json.dumps(receipt, indent=2)+"\n")
    print(f"Saved: {OUT.relative_to(ROOT)}")
    print(f"Validated: {len(table)} rows, 8 columns; source values, spending identity and Excel read-back.")


if __name__ == "__main__":
    main()
